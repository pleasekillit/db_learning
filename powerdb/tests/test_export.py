"""Excel 导出测试：验证业务视图字段与中文表头一致。"""
from datetime import datetime

import pytest
from openpyxl import load_workbook

from powerdb.config import DatabaseConfig
from powerdb.core.engine import create_schema, init_engine, session_scope
from powerdb.exporters import (
    CASE_RECORD_FIELDS,
    export_case_records_to_excel,
)
from powerdb.models import (
    DataFile,
    Project,
    Sample,
    Standard,
    StandardItem,
    TestRecord,
    TestRun,
)


@pytest.fixture()
def seeded(tmp_path):
    init_engine(DatabaseConfig(url=f"sqlite:///{tmp_path / 'test.db'}"))
    create_schema(drop=True)

    with session_scope() as session:
        project = Project(name="P")
        sample = Sample(code="S1", name="样品1", project=project)
        std = Standard(code="GB", name="标准")
        item = StandardItem(code="EFF", name="额定效率", unit="%", standard=std)
        session.add_all([project, sample, std, item])
        session.flush()

        run = TestRun(sample_id=sample.id, run_no="R1")
        session.add(run)
        session.flush()

        rec = TestRecord(
            test_run_id=run.id, standard_item_id=item.id, status="pass",
            started_at=datetime(2026, 9, 26, 10, 0, 0),
        )
        session.add(rec)
        session.flush()

        session.add_all([
            DataFile(test_record_id=rec.id, name="a.raw", path="/x/a.raw", kind="raw"),
            DataFile(test_record_id=rec.id, name="b.report", path="/x/b.report", kind="report"),
        ])


def test_export_headers_and_row(seeded, tmp_path):
    out = tmp_path / "out.xlsx"
    path = export_case_records_to_excel(out)

    assert path == out
    assert out.exists()

    wb = load_workbook(out)
    ws = wb.active

    # 表头 = 字段映射表里定义的中文列名，顺序一致
    headers = [cell.value for cell in ws[1]]
    assert headers == [f.header for f in CASE_RECORD_FIELDS]

    # 第二行是那条用例执行记录
    row = [cell.value for cell in ws[2]]
    assert row[headers.index("用例编号")] == "EFF"
    assert row[headers.index("用例名称")] == "额定效率"
    assert row[headers.index("样品编号")] == "S1"
    assert row[headers.index("试验编号")] == "R1"
    assert row[headers.index("执行时间")] == "2026-09-26 10:00:00"
    assert row[headers.index("执行结果")] == "通过"  # pass -> 中文
    assert row[headers.index("相关数据文件")] == "a.raw; b.report"